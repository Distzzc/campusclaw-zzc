import pytest
from knowledge.text import extract, split_units, ExtractionError


class CharacterTokenizer:
    def encode(self, text, add_special_tokens=False):
        return list(text)

    def num_special_tokens_to_add(self, pair=False):
        return 2


def test_chunking_preserves_long_chinese_text_and_locations():
    original = '植物通过光合作用制造养分。' * 100
    units = [(original, {'page': 3})]
    first = split_units(units, CharacterTokenizer(), 32, 1, 'test')
    assert first == split_units(units, CharacterTokenizer(), 32, 1, 'test')
    covered = set()
    for chunk in first[1]:
        loc = chunk['location']
        assert loc['page'] == 3
        assert chunk['text'] == original[loc['char_start']:loc['char_end']]
        assert len(chunk['text']) <= 30
        covered.update(range(loc['char_start'], loc['char_end']))
    assert covered == set(range(len(original)))
    assert first[0] != split_units(units, CharacterTokenizer(), 32, 1, 'new-model')[0]


def test_text_markdown_and_invalid_formats(tmp_path):
    path = tmp_path / 'a.md'
    path.write_text('# 课程\n\n光合作用\n\n## 实验\n\n观察叶片', encoding='utf8')
    units = extract(path, path.name)
    assert units[-1][1] == {'paragraph': 4, 'headings': ['课程', '实验']}
    with pytest.raises(ExtractionError):
        extract(path, 'a.exe')
    path.write_text('   ')
    with pytest.raises(ExtractionError):
        extract(path, 'a.txt')


def test_docx_paragraph_and_table(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_heading('课程', 1)
    doc.add_paragraph('植物通过光合作用制造养分')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格知识'
    path = tmp_path / 'a.docx'
    doc.save(path)
    units = extract(path, path.name)
    assert units[1][1]['headings'] == ['课程']
    assert units[-1][0] == '表格知识'


def test_pdf_page_and_empty_scanned_page(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 10 100 Td (Photosynthesis lesson) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    path = tmp_path / 'a.pdf'
    writer.write(path)
    assert extract(path, path.name)[0] == ('Photosynthesis lesson', {'page': 1})
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(path)
    with pytest.raises(ExtractionError):
        extract(path, path.name)


def test_chinese_and_english_terms():
    from knowledge.text import terms
    assert '光合作用' in terms('植物通过光合作用制造养分')
    assert terms('HELLO hello') == ['hello']
    assert terms('?!"') == []
